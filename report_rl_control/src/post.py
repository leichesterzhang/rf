import copy
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
d = Document('raw.docx')
M = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
n = 0
for p in d.paragraphs:
    paras = p._p.findall(f'{{{M}}}oMathPara')
    if not paras:
        continue
    n += 1
    omp = paras[0]
    omath = omp.find(f'{{{M}}}oMath')
    # remove all children except pPr
    for ch in list(p._p):
        if ch.tag != qn('w:pPr'):
            p._p.remove(ch)
    pf = p.paragraph_format
    pf.first_line_indent = Pt(0); pf.left_indent = Pt(0)
    pf.space_before = Pt(4); pf.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    tabs = OxmlElement('w:tabs')
    for val, pos in (('center', 4535), ('right', 9070)):
        t = OxmlElement('w:tab'); t.set(qn('w:val'), val); t.set(qn('w:pos'), str(pos)); tabs.append(t)
    pPr.append(tabs)
    def run_tab():
        r = OxmlElement('w:r'); r.append(OxmlElement('w:tab')); return r
    p._p.append(run_tab()); p._p.append(omath); p._p.append(run_tab())
    r = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = f'({n})'; r.append(t); p._p.append(r)
print('equations', n)

def border(tag, sz):
    e = OxmlElement(f'w:{tag}'); e.set(qn('w:val'), 'single' if sz else 'nil')
    if sz: e.set(qn('w:sz'), str(sz)); e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000')
    return e
for tb in d.tables:
    tblPr = tb._tbl.tblPr
    for old in tblPr.findall(qn('w:tblBorders')) + tblPr.findall(qn('w:tblStyle')):
        tblPr.remove(old)
    b = OxmlElement('w:tblBorders')
    for tag, sz in (('top', 12), ('left', 0), ('bottom', 12), ('right', 0), ('insideH', 0), ('insideV', 0)):
        b.append(border(tag, sz))
    tblPr.append(b)
    jc = OxmlElement('w:jc'); jc.set(qn('w:val'), 'center'); tblPr.append(jc)
    w = tblPr.find(qn('w:tblW'))
    if w is None:
        w = OxmlElement('w:tblW'); tblPr.append(w)
    w.set(qn('w:type'), 'pct'); w.set(qn('w:w'), '5000')
    for cell in tb.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tb_ = OxmlElement('w:tcBorders'); tb_.append(border('bottom', 6)); tcPr.append(tb_)
        for par in cell.paragraphs:
            for r in par.runs: r.bold = True
    for row in tb.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.first_line_indent = Pt(0)
                for r in par.runs: r.font.size = Pt(10.5)
from docx.text.paragraph import Paragraph
for tb in d.tables:
    nxt = tb._tbl.getnext()
    if nxt is not None and nxt.tag == qn('w:p'):
        Paragraph(nxt, None).paragraph_format.space_before = Pt(9)
# references: no first-line indent, hanging
start = False
for p in d.paragraphs:
    if p.style.name.startswith('Heading') and p.text.strip() == '参考文献':
        start = True; continue
    if start:
        pf = p.paragraph_format; pf.first_line_indent = Pt(-21); pf.left_indent = Pt(21)
        for r in p.runs: r.font.size = Pt(10.5)
d.save('强化学习运动控制_技术报告.docx')
